"""文档加载：把 txt / pdf / docx / csv 读取为 LangChain Document。"""
import os

import pandas as pd
from langchain_core.documents import Document

from app.core.progress import fmt_size, info


def load_single_file(file_path) -> Document | None:
    """读取单个文件；不支持的类型或读取失败返回 None。

    返回的 Document.metadata 只带 file（文件名），后续入库时透传到每个子块。
    """
    file_path = str(file_path)
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext == ".txt":
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
        elif ext == ".pdf":
            import fitz

            doc = fitz.open(file_path)
            text = "\n".join(page.get_text() for page in doc)
        elif ext == ".docx":
            import docx

            text = "\n".join(p.text for p in docx.Document(file_path).paragraphs)
        elif ext == ".csv":
            text = pd.read_csv(file_path).to_string(index=False)
        else:
            return None
    except Exception as e:
        print(f"读取失败 {file_path}: {e}")
        return None
    info(
        f"读取文件：{os.path.basename(file_path)}（{ext}，"
        f"{fmt_size(os.path.getsize(file_path))}，{len(text)} 字）"
    )
    return Document(page_content=text, metadata={"file": os.path.basename(file_path)})
